"""Data layer for the single-sheet CN rail workbook.

The repo version reads one worksheet (default ``Shipments``) containing one row
per container: shipment details, vessel ETA, discharge date, carrier rail ETA,
and the actual rail events (RAMPED / AT DESTINATION / DERAMPED / ...).

The workbook is not committed to git. Create an empty template with
``python scripts/make_template.py --out data/CN_Rail_Data.xlsx`` and fill it in.
"""
from __future__ import annotations

import csv
import json
import os
import re
from collections import defaultdict
from datetime import date, datetime
from zoneinfo import ZoneInfo

import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE_DIR, 'data')
DATA_XLSX = os.path.join(DATA_DIR, 'CN_Rail_Data.xlsx')
DATA_CSV = os.path.join(DATA_DIR, 'CN_Rail_Data.csv')
CACHE_DIR = os.path.join(BASE_DIR, 'cache')
MASTER_CSV = os.path.join(CACHE_DIR, 'master.csv')
MASTER_META = os.path.join(CACHE_DIR, 'master.json')
SHEET_NAME = 'Shipments'

TZ = {'PT': 'America/Vancouver', 'MT': 'America/Edmonton', 'CT': 'America/Winnipeg',
      'ET': 'America/Toronto', 'AT': 'America/Halifax'}

COLUMNS = (
    ('container', 'Container Number'),
    ('job', 'Job#'),
    ('client', 'Client'),
    ('consignee', 'Consignee'),
    ('shipping_line', 'Shipping Line'),
    ('pod', 'Port of Discharge'),
    ('ramp_station', 'Origin Ramp'),
    ('rail_dest', 'Final Destination'),
    ('port_eta', 'Vessel ETA'),
    ('discharge_date', 'Discharge Date'),
    ('tracker_rail_eta', 'Rail ETA'),
    ('ramp_ts', 'RAMPED'),
    ('first_dep_ts', 'DEPARTED'),
    ('at_dest_ts', 'AT DESTINATION'),
    ('deramp_ts', 'DERAMPED'),
    ('outgate_ts', 'OUT-GATE'),
    ('delivered_ts', 'DELIVERED'),
    ('notes', 'Notes'),
)
REQUIRED_COLUMNS = ('container', 'pod', 'ramp_station', 'rail_dest', 'port_eta')
DATE_COLUMNS = ('port_eta', 'discharge_date', 'tracker_rail_eta', 'ramp_ts',
                'first_dep_ts', 'at_dest_ts', 'deramp_ts', 'outgate_ts', 'delivered_ts')

_ALIASES = {
    'container number': 'container', 'container': 'container',
    'container no': 'container', 'equipmentid': 'container',
    'equipment id': 'container', 'equipment': 'container',
    'job#': 'job', 'job': 'job', 'job number': 'job',
    'client': 'client', 'consignee': 'consignee',
    'shipping line': 'shipping_line', 'carrier': 'shipping_line',
    'port of discharge': 'pod', 'pod': 'pod', 'port': 'pod',
    'origin ramp': 'ramp_station', 'ramp': 'ramp_station',
    'rail origin': 'ramp_station', 'ramp station': 'ramp_station',
    'final destination': 'rail_dest', 'destination': 'rail_dest',
    'rail destination': 'rail_dest',
    'vessel eta': 'port_eta', 'port eta': 'port_eta', 'eta date': 'port_eta',
    'eta': 'port_eta', 'vessel arrival': 'port_eta',
    'discharge date': 'discharge_date', 'actual discharge': 'discharge_date',
    'vessel discharge': 'discharge_date', 'discharged': 'discharge_date',
    'rail eta': 'tracker_rail_eta', 'carrier rail eta': 'tracker_rail_eta',
    'tracker rail eta': 'tracker_rail_eta',
    'ramped': 'ramp_ts', 'loaded on rail': 'ramp_ts', 'rail loaded': 'ramp_ts',
    'departed': 'first_dep_ts', 'rail departure': 'first_dep_ts',
    'at destination': 'at_dest_ts', 'arrival': 'at_dest_ts',
    'rail arrival': 'at_dest_ts',
    'deramped': 'deramp_ts', 'empty': 'deramp_ts', 'deramp': 'deramp_ts',
    'out-gate': 'outgate_ts', 'outgate': 'outgate_ts', 'out gate': 'outgate_ts',
    'delivered': 'delivered_ts', 'delivery': 'delivered_ts',
    'pad placement': 'delivered_ts', 'at receiver': 'delivered_ts',
    'notes': 'notes', 'note': 'notes', 'comment': 'notes',
}


def _key(text):
    return re.sub(r'\s+', ' ', str(text or '').strip().lower().replace('_', ' '))


def _alias_map():
    return {_key(alias): key for alias, key in _ALIASES.items()}


def data_path(path=None):
    if path:
        return path
    env = os.environ.get('CN_RAIL_DATA')
    if env:
        return env
    if os.path.exists(DATA_XLSX):
        return DATA_XLSX
    if os.path.exists(DATA_CSV):
        return DATA_CSV
    return DATA_XLSX


def parse_ts(s):
    m = re.match(r'(\d{4}-\d{2}-\d{2}) (\d{2}:\d{2}) ([A-Z]{1,4})', str(s))
    if not m:
        return None
    dt = datetime.strptime(m.group(1) + ' ' + m.group(2), '%Y-%m-%d %H:%M')
    tz = TZ.get(m.group(3))
    return dt.replace(tzinfo=ZoneInfo(tz)) if tz else None


def parse_dt(s):
    if s is None:
        return None
    if isinstance(s, datetime):
        return s
    if isinstance(s, date):
        return datetime.combine(s, datetime.min.time())
    s = str(s).strip()
    if not s or s.upper().startswith('TBA'):
        return None
    tagged = parse_ts(s)
    if tagged is not None:
        return tagged
    try:
        return datetime.fromisoformat(s.replace('Z', '+00:00'))
    except ValueError:
        pass
    m = re.match(r'(\d{4}-\d{2}-\d{2})(?:[ T](\d{2}:\d{2}))?', s)
    if m:
        return datetime.strptime(m.group(1) + ' ' + (m.group(2) or '00:00'), '%Y-%m-%d %H:%M')
    for fmt in ('%d/%m/%Y', '%d/%m/%Y %H:%M', '%m/%d/%Y', '%m/%d/%Y %H:%M'):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def parse_ddmmyyyy(s):
    try:
        return datetime.strptime(str(s).strip(), '%d/%m/%Y')
    except Exception:
        return None


_POD_ALIASES = {
    'MONTREAL': 'MONTREAL', 'MTL': 'MONTREAL',
    'VANCOUVER': 'VANCOUVER', 'VANC': 'VANCOUVER',
    'PRINCE RUPERT': 'PRINCE RUPERT', 'PR': 'PRINCE RUPERT',
    'HALIFAX': 'HALIFAX',
    'NEW YORK': 'NEW YORK', 'NEWARK': 'NEWARK', 'NY': 'NEW YORK',
    "SAINT JOHN'S": 'SAINT JOHN', 'SAINT JOHN': 'SAINT JOHN',
    'PHILADELPHIA': 'PHILADELPHIA',
    'CALGARY': 'CALGARY',
}


def canon_pod(v):
    if v is None:
        return None
    return _POD_ALIASES.get(str(v).strip().upper(), str(v).strip().upper())


PORT_RAMPS = {
    'MONTREAL': ['MONTREAL VIAU TERM', 'MONTREAL TERMONT', 'MONTREAL WHARF'],
    'VANCOUVER': ['ROBERTS BANK', 'VANCOUVER VANTERM', 'FRASER SURREY IMPEX', 'FAIRVIEW COVE IMPEX'],
    'PRINCE RUPERT': ['PRINCE RUPERT TERM'],
    'HALIFAX': ['HALIFAX HALTERM TER'],
}
RAMP_PORT = {r: p for p, ramps in PORT_RAMPS.items() for r in ramps}


def canon_dest(v):
    """Canonical rail destination terminal from any free-form destination string."""
    if v is None:
        return None
    x = str(v).split('\n')[0].strip().upper().replace('-OG', '').strip()
    if not x or x in ('NA', 'N/A', 'NONE'):
        return None
    if 'MALPORT' in x:
        return 'MALPORT'
    if 'BRAMPTON' in x:
        return 'BRAMPTON INTER TERM'
    if 'CALGARY' in x:
        return 'CALGARY LOG PARK'
    if 'EDMONTON' in x:
        return 'EDMONTON INTER TERM'
    if 'TASCH' in x:
        return 'MONTREAL TASCH YARD'
    if 'WINNIPEG' in x:
        return 'WINNIPEG'
    if 'MUSKET' in x:
        return 'MUSKET'
    if 'MONTREAL' in x:
        return 'MONTREAL'
    if 'TORONTO' in x or 'CP TORONTO' in x:
        return 'MALPORT'
    if 'DETROIT' in x:
        return 'DETROIT INTER TERM'
    if 'ROBERTS BANK' in x:
        return 'ROBERTS BANK'
    if 'PRINCE RUPERT' in x:
        return 'PRINCE RUPERT TERM'
    if 'REGINA' in x:
        return 'REGINA INTER TERM'
    if 'MONCTON' in x:
        return 'MONCTON INTER TERM'
    return x


def _rows_from_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]
        raw = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not raw:
        return [], []
    headers = [str(h).strip() if h is not None else '' for h in raw[0]]
    return headers, [r for r in raw[1:] if any(v not in (None, '') for v in r)]


def _rows_from_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        raw = list(reader)
    if not raw:
        return [], []
    headers = [str(h).strip() for h in raw[0]]
    return headers, raw[1:]


def _read_raw(path):
    if str(path).lower().endswith('.csv'):
        return _rows_from_csv(path)
    return _rows_from_workbook(path)


def _header_index(headers):
    aliases = _alias_map()
    out = {}
    for i, h in enumerate(headers):
        key = aliases.get(_key(h))
        if key and key not in out:
            out[key] = i
    return out


def read_records(path=None):
    """Read the single sheet, returning normalized row dictionaries."""
    path = data_path(path)
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"Data workbook not found: {path}\n"
            "Create one with scripts/make_template.py or select a workbook in the app.")
    headers, raw = _read_raw(path)
    idx = _header_index(headers)
    missing = [label for key, label in COLUMNS if key in REQUIRED_COLUMNS and key not in idx]
    if missing:
        raise ValueError('Missing required column(s): ' + ', '.join(missing))
    records = []
    for row_no, row in enumerate(raw, start=2):
        if not row or all(v in (None, '') for v in row):
            continue
        rec = {}
        for key, _label in COLUMNS:
            i = idx.get(key)
            rec[key] = row[i] if i is not None and i < len(row) else None
        container = str(rec['container'] or '').strip().upper()
        if not container:
            continue
        if '\n' in container or ',' in container:
            raise ValueError(f'Row {row_no}: put one container number per row.')
        rec['container'] = container
        for key in ('job', 'client', 'consignee', 'shipping_line', 'notes'):
            rec[key] = (str(rec[key]).strip() if rec[key] not in (None, '') else None)
        rec['pod'] = canon_pod(rec['pod'])
        rec['rail_dest'] = canon_dest(rec['rail_dest'])
        rec['ramp_station'] = (str(rec['ramp_station']).strip().upper()
                               if rec['ramp_station'] not in (None, '') else None)
        for key in DATE_COLUMNS:
            rec[key] = parse_dt(rec[key])
        records.append(rec)
    if not records:
        raise ValueError('The sheet has no data rows.')
    return records


def _gap(row, a, b):
    if not row.get(a) or not row.get(b):
        return None
    x, y = row[a], row[b]
    if isinstance(x, datetime):
        x = x.replace(tzinfo=None)
    if isinstance(y, datetime):
        y = y.replace(tzinfo=None)
    return round((x - y).total_seconds() / 86400.0, 3)


def _cache_valid(path):
    if not (os.path.exists(MASTER_CSV) and os.path.exists(MASTER_META)):
        return False
    try:
        with open(MASTER_META, encoding='utf-8') as f:
            meta = json.load(f)
        st = os.stat(path)
        return (meta.get('source') == os.path.abspath(path)
                and meta.get('mtime_ns') == st.st_mtime_ns
                and meta.get('size') == st.st_size)
    except (OSError, ValueError):
        return False


def build_master(refresh=False, path=None):
    """Build one training row per container that has a RAMPED event."""
    path = data_path(path)
    if not refresh and _cache_valid(path):
        return load_master_csv()
    records = read_records(path)
    rows = []
    for rec in records:
        if not rec['ramp_ts']:
            continue
        row = {key: rec.get(key) for key, _label in COLUMNS}
        row['dwell_days'] = _gap(rec, 'ramp_ts', 'port_eta')
        row['transit_days'] = _gap(rec, 'at_dest_ts', 'ramp_ts')
        row['deramp_days'] = _gap(rec, 'deramp_ts', 'ramp_ts')
        row['outgate_days'] = _gap(rec, 'outgate_ts', 'ramp_ts')
        row['total_days'] = _gap(rec, 'at_dest_ts', 'port_eta')
        row['total_deramp_days'] = _gap(rec, 'deramp_ts', 'port_eta')
        rows.append(row)
    if not rows:
        raise ValueError('No rows contain a RAMPED timestamp, so there is nothing to train on.')
    os.makedirs(CACHE_DIR, exist_ok=True)
    fields = [key for key, _ in COLUMNS] + [
        'dwell_days', 'transit_days', 'deramp_days', 'outgate_days',
        'total_days', 'total_deramp_days']
    with open(MASTER_CSV, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for r in rows:
            w.writerow({k: (v.isoformat() if isinstance(v, datetime) else v) for k, v in r.items()})
    st = os.stat(path)
    with open(MASTER_META, 'w', encoding='utf-8') as f:
        json.dump({'source': os.path.abspath(path), 'mtime_ns': st.st_mtime_ns,
                   'size': st.st_size}, f)
    return rows


def load_master_csv():
    rows = []
    with open(MASTER_CSV, newline='', encoding='utf-8') as f:
        for r in csv.DictReader(f):
            for k in DATE_COLUMNS:
                if r.get(k):
                    try:
                        r[k] = datetime.fromisoformat(r[k])
                    except ValueError:
                        r[k] = None
                else:
                    r[k] = None
            for k in ('dwell_days', 'transit_days', 'deramp_days', 'outgate_days',
                      'total_days', 'total_deramp_days'):
                if r.get(k) not in (None, ''):
                    r[k] = float(r[k])
                else:
                    r[k] = None
            rows.append(r)
    return rows


def all_known_containers(path=None):
    """Every container in the sheet, including shipments with no rail events."""
    out = {}
    for rec in read_records(path):
        out[rec['container']] = dict(rec)
    return out


def validate_data_file(path):
    """Return (ok, message) for the app's upload dialog."""
    if not path:
        return False, 'No file selected.'
    if not os.path.exists(path):
        return False, 'File not found.'
    if not str(path).lower().endswith(('.xlsx', '.xlsm', '.csv')):
        return False, 'Choose an .xlsx, .xlsm or .csv file.'
    try:
        records = read_records(path)
    except Exception as exc:
        return False, str(exc)
    with_ramp = sum(1 for r in records if r['ramp_ts'])
    with_eta = sum(1 for r in records if r['port_eta'])
    if not with_ramp:
        return False, 'No RAMPED timestamps found; at least one completed journey is needed.'
    if not with_eta:
        return False, 'No vessel ETA values found; the model needs an anchor date.'
    return True, (f'{len(records)} containers, {with_ramp} with RAMPED, '
                  f'{with_eta} with vessel ETA.')


if __name__ == '__main__':
    import statistics
    rows = build_master(refresh=True)
    print('master rows:', len(rows))
    with_eta = [r for r in rows if r['port_eta']]
    print('with port ETA:', len(with_eta))
    print('with port ETA + arrival:', sum(1 for r in with_eta if r['at_dest_ts']))
    print()
    print('--- dwell (vessel ETA -> RAMPED) by POD, days ---')
    byp = defaultdict(list)
    for r in with_eta:
        if r['dwell_days'] is not None:
            byp[r['pod']].append(r['dwell_days'])
    for p, v in sorted(byp.items(), key=lambda kv: -len(kv[1])):
        s = sorted(v)
        print(f'  {p:14s} n={len(s):3d} p10={s[len(s)//10]:5.1f} '
              f'p50={statistics.median(s):5.1f} p90={s[len(s)*9//10]:5.1f}')
    print()
    print('--- transit (RAMPED -> AT DEST) by lane, days ---')
    byl = defaultdict(list)
    for r in rows:
        if r['transit_days'] is not None and r['rail_dest']:
            byl[(r['ramp_station'], r['rail_dest'])].append(r['transit_days'])
    for k, v in sorted(byl.items(), key=lambda kv: -len(kv[1])):
        if len(v) < 5:
            continue
        s = sorted(v)
        print(f'  {len(s):4d}  {k[0]:22s} -> {k[1]:20s} '
              f'p10={s[len(s)//10]:5.1f} p50={statistics.median(s):5.1f} '
              f'p90={s[len(s)*9//10]:5.1f}')
