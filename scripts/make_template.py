"""Create the single-sheet workbook template used by the repo version.

Examples:
    python scripts/make_template.py --out data/CN_Rail_Data.xlsx
    python scripts/make_template.py --demo --out data/CN_Rail_Data.xlsx
"""
from __future__ import annotations

import argparse
import os
import random
import sys
from datetime import datetime, timedelta

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from cn_model.data import PORT_RAMPS, SHEET_NAME  # noqa: E402

HEADERS = [
    'Container Number', 'Job#', 'Client', 'Consignee', 'Shipping Line',
    'Port of Discharge', 'Origin Ramp', 'Final Destination',
    'Vessel ETA', 'Discharge Date', 'Rail ETA',
    'RAMPED', 'DEPARTED', 'AT DESTINATION', 'DERAMPED', 'OUT-GATE', 'DELIVERED',
    'Notes',
]
DATE_HEADERS = {'Vessel ETA', 'Discharge Date', 'Rail ETA', 'RAMPED', 'DEPARTED',
                'AT DESTINATION', 'DERAMPED', 'OUT-GATE', 'DELIVERED'}
WIDTHS = [18, 12, 18, 24, 16, 18, 24, 26, 18, 18, 18, 18, 18, 18, 18, 18, 18, 30]
COMMENTS = {
    'Container Number': 'One container per row. Required.',
    'Port of Discharge': 'Vessel discharge port, e.g. MONTREAL or VANCOUVER.',
    'Origin Ramp': 'CN rail ramp where the container is loaded.',
    'Final Destination': 'Rail destination terminal.',
    'Vessel ETA': 'Prediction anchor; vessel ETA at the port.',
    'Discharge Date': 'Actual vessel discharge date, when known.',
    'Rail ETA': 'Carrier/tracker rail ETA, when available.',
    'RAMPED': 'Actual loaded-on-rail timestamp.',
    'AT DESTINATION': 'Actual rail arrival timestamp.',
    'DERAMPED': 'Actual empty/deramp timestamp.',
}

DESTINATIONS = ['BRAMPTON INTER TERM', 'CALGARY LOG PARK', 'EDMONTON INTER TERM',
                'MALPORT', 'MONTREAL TASCH YARD', 'WINNIPEG']
DWELL_BASE = {'MONTREAL': 7.0, 'VANCOUVER': 8.0, 'PRINCE RUPERT': 9.0,
              'HALIFAX': 6.0, 'NEW YORK': 5.5}


def _build_rows(count=130, seed=7):
    rng = random.Random(seed)
    pods = [p for p in DWELL_BASE if PORT_RAMPS.get(p)]
    rows = []
    start = datetime(2025, 6, 1, 8, 0)
    span = (datetime(2026, 9, 28) - start).days
    for i in range(count):
        eta = start + timedelta(days=rng.randrange(span), hours=rng.randrange(0, 8))
        pod = rng.choice(pods)
        ramp = rng.choice(PORT_RAMPS[pod])
        dest = rng.choice(DESTINATIONS)
        dwell = max(1.0, rng.gauss(DWELL_BASE[pod], 3.0))
        ramp_ts = eta + timedelta(days=dwell)
        transit = max(1.0, rng.gauss(6.0, 3.0))
        arrival = ramp_ts + timedelta(days=transit)
        deramp = arrival + timedelta(hours=rng.choice([2, 4, 8, 20]))
        departed = ramp_ts + timedelta(hours=rng.randrange(1, 12))
        outgate = deramp + timedelta(days=rng.choice([0, 1, 2]))
        discharge = eta + timedelta(days=rng.choice([0, 0, 0, 1, 3]))
        carrier = arrival + timedelta(days=rng.gauss(1.0, 1.2))
        if i >= count - 8:
            arrival = deramp = outgate = None
        rows.append({
            'Container Number': f'DEMU{1000000 + i * 137:07d}',
            'Job#': f'DEMO{1000 + i}',
            'Client': rng.choice(['Northline', 'Maple Freight', 'Prairie Goods']),
            'Consignee': rng.choice(['Consignee A', 'Consignee B', 'Consignee C']),
            'Shipping Line': rng.choice(['Maersk', 'ONE', 'COSCO', 'CMA CGM']),
            'Port of Discharge': pod,
            'Origin Ramp': ramp,
            'Final Destination': dest,
            'Vessel ETA': eta,
            'Discharge Date': discharge,
            'Rail ETA': carrier,
            'RAMPED': ramp_ts,
            'DEPARTED': departed,
            'AT DESTINATION': arrival,
            'DERAMPED': deramp,
            'OUT-GATE': outgate,
            'DELIVERED': outgate,
            'Notes': 'Synthetic demo row; not real shipment data.',
        })
    return rows


def create_workbook(path, demo=False, rows=130):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    head_fill = PatternFill('solid', fgColor='1F4E79')
    head_font = Font(color='FFFFFF', bold=True)
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.comment = Comment(COMMENTS.get(header, header), 'CN Rail model')
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS[col - 1]
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(HEADERS))}1'

    if demo:
        for r, data in enumerate(_build_rows(rows), start=2):
            for c, header in enumerate(HEADERS, start=1):
                cell = ws.cell(row=r, column=c, value=data.get(header))
                if header in DATE_HEADERS:
                    cell.number_format = 'yyyy-mm-dd hh:mm'

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--out', default=os.path.join(ROOT, 'data', 'CN_Rail_Data.xlsx'))
    parser.add_argument('--demo', action='store_true', help='fill with synthetic rows')
    parser.add_argument('--rows', type=int, default=130)
    args = parser.parse_args()
    path = create_workbook(args.out, demo=args.demo, rows=args.rows)
    kind = 'demo workbook' if args.demo else 'empty template'
    print(f'Created {kind}: {path}')


if __name__ == '__main__':
    main()
